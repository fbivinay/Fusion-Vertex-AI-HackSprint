import csv
import io
import json
import mimetypes
import threading
import time
from collections import defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from google import genai
from google.genai import types
import pymupdf as fitz

from .config import get_settings
from .repository import repository
from .schemas import AnalysisResult, AskResult, BRDNarrative
from .storage import load_source, save_output, save_processed


settings = get_settings()

_genai_client: genai.Client | None = None
_request_lock = threading.Lock()
_last_request_started = 0.0


def _client() -> genai.Client:
    global _genai_client
    if _genai_client is None:
        if not settings.gemini_api_key:
            raise RuntimeError("Set GEMINI_API_KEY to a Google AI Studio API key before requesting Gemini analysis.")
        _genai_client = genai.Client(api_key=settings.gemini_api_key)
    return _genai_client


def _extract_text(filename: str, mime_type: str, content: bytes) -> tuple[str, list[tuple[str, bytes, str]]]:
    extension = Path(filename).suffix.lower()
    attachments: list[tuple[str, bytes, str]] = []
    if extension == ".pdf" or mime_type == "application/pdf":
        document = fitz.open(stream=content, filetype="pdf")
        pages = [f"[Page {index + 1}]\n{page.get_text('text')}" for index, page in enumerate(document)]
        if not any(page.strip() for page in pages):
            attachments.append((filename, content, "application/pdf"))
        return "\n\n".join(pages), attachments
    if extension == ".docx" or mime_type.endswith("wordprocessingml.document"):
        document = Document(io.BytesIO(content))
        lines: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                lines.append(paragraph.text.strip())
        for table_index, table in enumerate(document.tables, start=1):
            for row in table.rows:
                lines.append(f"[Table {table_index}] " + " | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(lines), attachments
    if extension == ".csv" or mime_type in {"text/csv", "application/csv"}:
        try:
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
            return "\n".join(" | ".join(row) for row in rows[:500]), attachments
        except UnicodeDecodeError:
            return content.decode("latin-1"), attachments
    if extension in {".png", ".jpg", ".jpeg"} or mime_type.startswith("image/"):
        detected = mimetypes.guess_type(filename)[0] or mime_type or "image/png"
        attachments.append((filename, content, detected))
        return "[Image source; inspect the attached visual. Cite location as the visible screen area or image region.]", attachments
    return content.decode("utf-8-sig", errors="replace"), attachments


def _read_all_sources(project_id: str) -> tuple[list[dict[str, Any]], list[types.Part]]:
    sources = repository.list("sources", project_id)
    if not sources:
        raise ValueError("Upload at least one source before running analysis.")
    prompt_sections: list[str] = []
    multimodal_parts: list[types.Part] = []
    total_bytes = 0
    valid_ids = {source["source_id"] for source in sources}
    for source in sources:
        content = load_source(source["object_uri"])
        total_bytes += len(content)
        if total_bytes > 20 * 1024 * 1024:
            raise ValueError("The combined source set is over 20 MB. Upload a smaller set and retry.")
        text, attachments = _extract_text(source["filename"], source["mime_type"], content)
        prompt_sections.append(
            f"SOURCE ID: {source['source_id']}\nFILE: {source['filename']}\nTYPE: {source['mime_type']}\nEXTRACTED CONTENT:\n{text[:50000]}"
        )
        for filename, data, mime in attachments:
            prompt_sections.append(f"VISUAL ATTACHMENT: {source['source_id']} — {filename} ({mime})")
            multimodal_parts.append(types.Part.from_bytes(data=data, mime_type=mime))
    base_prompt = f"""You are Fusion Vertex AI, a business requirements analysis system. Analyze every supplied source TOGETHER as one project context. Do not treat sources independently. Separate source facts from inferences, never resolve contradictions silently, and never invent evidence.

Return a structured business analysis with objectives, problem, stakeholders, assumptions, missing information, risks, requirements, and conflicts. For each requirement:
- use a stable ID such as FR-001, NFR-001, or BR-001; type must match the prefix;
- classification must be SUPPORTED when directly evidenced, INFERRED when reasonable but not directly stated, CONFLICTED when sources disagree, or NEEDS_HUMAN_DECISION when a human must choose;
- cite one or more exact source IDs from the source inventory, plus page/section/row/screen area where available, a short supporting quote or visual fact, and a concise reason;
- give a confidence from 0 to 1 and priority P0-P3; do not overstate certainty;
- include testable acceptance criteria.

For conflicts, cite two different source IDs and quote each side, explain business and technical impact, explain why it matters, and associate affected requirement IDs. Keep the status NEEDS_HUMAN_DECISION. Never choose a winning source.

Use concise evidence-based explanations only. Do not expose hidden chain-of-thought. Distinguish missing information from assumptions. Source IDs must exactly match the inventory.

SOURCE INVENTORY AND EXTRACTED TEXT:\n{chr(10).join(prompt_sections)}"""
    multimodal_parts.insert(0, types.Part.from_text(text=base_prompt))
    return sources, multimodal_parts


def _generate_json(prompt_parts: list[types.Part], schema: type[Any]) -> Any:
    global _last_request_started
    with _request_lock:
        wait_for = settings.gemini_min_interval_seconds - (time.monotonic() - _last_request_started)
        if _last_request_started and wait_for > 0:
            time.sleep(wait_for)
        _last_request_started = time.monotonic()
        response = _client().models.generate_content(
            model=settings.gemini_model,
            contents=[types.Content(role="user", parts=prompt_parts)],
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=schema,
                thinking_config=types.ThinkingConfig(thinking_level="LOW"),
                max_output_tokens=8192,
            ),
        )
    if not response.text:
        raise RuntimeError("Gemini returned an empty structured response.")
    return schema.model_validate_json(response.text)


def analyze_project(project_id: str) -> dict[str, Any]:
    sources, parts = _read_all_sources(project_id)
    analysis: AnalysisResult = _generate_json(parts, AnalysisResult)
    source_ids = {source["source_id"] for source in sources}
    counts: defaultdict[str, int] = defaultdict(int)
    valid_requirement_ids: set[str] = set()
    evidence_rows: list[dict[str, Any]] = []
    provisional_to_final: dict[str, str] = {}
    validated_requirements = []
    for requirement in analysis.requirements:
        old_id = requirement.id
        evidence_row_start = len(evidence_rows)
        cleaned_refs = []
        for ref in requirement.evidence:
            if ref.source_id not in source_ids:
                continue
            cleaned_refs.append(ref)
            evidence_id = f"E-{len(evidence_rows) + 1:03d}"
            evidence_rows.append({
                "evidence_id": evidence_id,
                "requirement_id": "",
                "source_id": ref.source_id,
                "location": ref.location,
                "supporting_text": ref.supporting_text,
                "explanation": ref.explanation,
            })
        if not cleaned_refs:
            analysis.open_questions.append(
                f"Potential requirement excluded because Gemini returned no verifiable source citation: {requirement.description}"
            )
            continue
        prefix = {"functional": "FR", "non-functional": "NFR", "business-rule": "BR"}[requirement.type]
        counts[prefix] += 1
        expected_id = f"{prefix}-{counts[prefix]:03d}"
        requirement.id = expected_id
        valid_requirement_ids.add(expected_id)
        if old_id:
            provisional_to_final[old_id] = expected_id
        requirement.evidence = cleaned_refs
        validated_requirements.append(requirement)
        for row in evidence_rows[evidence_row_start:]:
            row["requirement_id"] = expected_id
    analysis.requirements = validated_requirements

    conflict_rows: list[dict[str, Any]] = []
    for index, conflict in enumerate(analysis.conflicts, start=1):
        conflict.id = f"C-{index:03d}"
        if conflict.source_a_id not in source_ids or conflict.source_b_id not in source_ids:
            continue
        mapped_ids = [provisional_to_final.get(rid, rid) for rid in conflict.requirement_ids]
        conflict.requirement_ids = [rid for rid in mapped_ids if rid in valid_requirement_ids]
        conflict.status = "NEEDS_HUMAN_DECISION"
        conflict_rows.append(conflict.model_dump())
        for requirement in analysis.requirements:
            if requirement.id in conflict.requirement_ids:
                requirement.classification = "NEEDS_HUMAN_DECISION"
                requirement.related_conflict_ids.append(conflict.id)

    # Reanalysis invalidates prior choices, but keeps the decision audit trail visible.
    prior_decisions = repository.list("decisions", project_id)
    for decision in prior_decisions:
        if decision.get("status") != "STALE":
            decision["status"] = "STALE"
            decision["stale_reason"] = "The source set was reanalyzed. Review this decision against the current conflict set."
            repository.put("decisions", project_id, decision["decision_id"], decision)
    # Replace analysis rows while preserving the human decision history.
    for entity in ("requirements", "evidence", "conflicts"):
        repository.delete_project_records(entity, project_id)
    for requirement in analysis.requirements:
        repository.put("requirements", project_id, requirement.id, requirement.model_dump())
    for row in evidence_rows:
        repository.put("evidence", project_id, row["evidence_id"], row)
    for conflict in conflict_rows:
        repository.put("conflicts", project_id, conflict["id"], conflict)
    for source in sources:
        source["analyzed"] = True
        repository.put("sources", project_id, source["source_id"], source)

    project = _get_project(project_id)
    project.update({
        "analysis_summary": analysis.concise_summary,
        "business_objectives": analysis.business_objectives,
        "problem_statement": analysis.problem_statement,
        "stakeholders": analysis.stakeholders,
        "assumptions": analysis.assumptions,
        "open_questions": analysis.open_questions,
        "risks": analysis.risks,
        "analysis_status": "complete",
    })
    repository.put("projects", project_id, project_id, project)
    result = {
        "project": project,
        "requirements": [item.model_dump() for item in analysis.requirements],
        "evidence": evidence_rows,
        "conflicts": conflict_rows,
        "decisions": repository.list("decisions", project_id),
    }
    save_processed(project_id, "analysis.json", json.dumps(result, indent=2, ensure_ascii=False).encode(), "application/json")
    return result


def _get_project(project_id: str) -> dict[str, Any]:
    for project in repository.list("projects", project_id):
        if project.get("project_id") == project_id:
            return project
    raise KeyError(project_id)


def answer_project_question(project_id: str, question: str) -> dict[str, Any]:
    project = _get_project(project_id)
    context = {
        "project": project,
        "requirements": repository.list("requirements", project_id),
        "evidence": repository.list("evidence", project_id),
        "conflicts": repository.list("conflicts", project_id),
        "decisions": repository.list("decisions", project_id),
        "sources": repository.list("sources", project_id),
    }
    prompt = f"""Answer the user's question using only this Fusion Vertex AI project record. Do not invent facts or answer generically. If the context lacks an answer, say what is missing. Cite relevant requirement, evidence, and source IDs in the response and structured citations. Keep the answer concise and do not reveal hidden chain-of-thought.

QUESTION: {question}
PROJECT CONTEXT JSON:\n{json.dumps(context, ensure_ascii=False, default=str)[:100000]}"""
    result: AskResult = _generate_json([types.Part.from_text(text=prompt)], AskResult)
    return result.model_dump()


def generate_brd(project_id: str) -> dict[str, Any]:
    project = _get_project(project_id)
    requirements = repository.list("requirements", project_id)
    if not requirements:
        raise ValueError("Run source analysis before generating the BRD.")
    conflicts = repository.list("conflicts", project_id)
    decisions = repository.list("decisions", project_id)
    sources = repository.list("sources", project_id)
    context = {
        "project": project,
        "requirements": requirements,
        "conflicts": conflicts,
        "decisions": decisions,
        "sources": [{key: source.get(key) for key in ("source_id", "filename", "mime_type")} for source in sources],
    }
    prompt = f"""Write polished, concise BRD narrative fields grounded only in this structured project analysis. Keep uncertainty visible, preserve every unresolved conflict and human decision, and never invent stakeholder approvals or business facts. User stories must be grounded in requirements; dependencies, risks, assumptions, and questions must match the analysis. Do not include hidden chain-of-thought.

PROJECT ANALYSIS JSON:\n{json.dumps(context, ensure_ascii=False, default=str)[:100000]}"""
    narrative: BRDNarrative = _generate_json([types.Part.from_text(text=prompt)], BRDNarrative)
    markdown = render_brd(project, narrative, requirements, conflicts, decisions, sources)
    uri = save_output(project_id, "business-requirements-document.md", markdown.encode("utf-8"), "text/markdown")
    project["brd_status"] = "generated"
    project["brd_uri"] = uri
    repository.put("projects", project_id, project_id, project)
    return {"project_id": project_id, "markdown": markdown, "narrative": narrative.model_dump(), "output_uri": uri}


def render_brd(project: dict[str, Any], narrative: BRDNarrative, requirements: list[dict[str, Any]], conflicts: list[dict[str, Any]], decisions: list[dict[str, Any]], sources: list[dict[str, Any]]) -> str:
    def bullets(items: list[str]) -> str:
        return "\n".join(f"- {item}" for item in items) if items else "- None identified in analyzed sources."

    lines = [
        f"# Business Requirements Document — {project['name']}",
        f"\n**Project ID:** `{project['project_id']}`  \n**Generated:** {project.get('updated_at', 'Analysis session')}  \n**Source count:** {len(sources)}",
        "\n## 1. Executive Summary\n" + narrative.executive_summary,
        "\n## 2. Business Objective\n" + narrative.business_objective,
        "\n## 3. Problem Statement\n" + narrative.problem_statement,
        "\n## 4. Stakeholders\n" + bullets(narrative.stakeholders or project.get("stakeholders", [])),
        "\n## 5. Proposed Solution\n" + narrative.proposed_solution,
        "\n## 6. Functional Requirements",
    ]
    for kind, title in (("functional", "Functional"), ("non-functional", "Non-functional"), ("business-rule", "Business rule")):
        lines.append(f"\n### {title}")
        group = [requirement for requirement in requirements if requirement["type"] == kind]
        if not group:
            lines.append("No requirements of this type were identified.")
        for requirement in group:
            lines.extend([
                f"\n#### {requirement['id']} — {requirement['description']}",
                f"**Priority:** {requirement['priority']} · **Confidence:** {requirement['confidence']:.0%} · **Conclusion:** {requirement['classification']}",
                f"\n{requirement.get('rationale') or 'No additional rationale provided.'}",
                "\nAcceptance criteria:\n" + bullets(requirement.get("acceptance_criteria", [])),
                "\nEvidence:\n" + ("\n".join(f"- {ref['source_id']} · {ref['location']}: “{ref['supporting_text']}” — {ref['explanation']}" for ref in requirement.get("evidence", [])) or "- No supporting source evidence returned."),
            ])
    lines.extend([
        "\n## 7. Business Rules\n" + bullets(narrative.business_rules),
        "\n## 8. User Stories\n" + bullets(narrative.user_stories),
        "\n## 9. Acceptance Criteria\n" + bullets([f"{req['id']}: {criterion}" for req in requirements for criterion in req.get("acceptance_criteria", [])]),
        "\n## 10. Dependencies\n" + bullets(narrative.dependencies),
        "\n## 11. Risks\n" + bullets(narrative.risks or project.get("risks", [])),
        "\n## 12. Assumptions\n" + bullets(narrative.assumptions or project.get("assumptions", [])),
        "\n## 13. Open Questions\n" + bullets(narrative.open_questions or project.get("open_questions", [])),
        "\n## 14. Conflicts",
    ])
    if conflicts:
        for conflict in conflicts:
            lines.append(f"\n### {conflict['id']} — {conflict['title']} · {conflict['severity']} · {conflict['status']}")
            lines.append(f"- **Source A ({conflict['source_a_id']}):** {conflict['source_a_statement']}")
            lines.append(f"- **Source B ({conflict['source_b_id']}):** {conflict['source_b_statement']}")
            lines.append(f"- **Impact:** {conflict['impact']}  \n- **Why it matters:** {conflict['why_it_matters']}")
    else:
        lines.append("\nNo conflicts were identified across the supplied sources.")
    final_decisions = [item for item in decisions if item.get("status") == "RECORDED"]
    decision_lines = [f"{item['conflict_id']}: {item['choice']} — {item.get('rationale', '')}" for item in decisions]
    decision_lines += [f"{item['id']}: pending human decision — {item['title']}" for item in conflicts if item['status'] == 'NEEDS_HUMAN_DECISION' and not any(d['conflict_id'] == item['id'] for d in final_decisions)]
    lines.extend(["\n## 15. Human Decisions", bullets(decision_lines), "\n## 16. Requirement Traceability Matrix", "\n| Requirement | Priority | Confidence | Conclusion | Source evidence |", "|---|---:|---:|---|---|"])
    for requirement in requirements:
        refs = "; ".join(f"{ref['source_id']} ({ref['location']})" for ref in requirement.get("evidence", [])) or "None"
        lines.append(f"| {requirement['id']} | {requirement['priority']} | {requirement['confidence']:.0%} | {requirement['classification']} | {refs} |")
    lines.append("\n---\nGenerated by Fusion Vertex AI from the listed source set. Statements marked inferred or unresolved require review.")
    return "\n".join(lines)

